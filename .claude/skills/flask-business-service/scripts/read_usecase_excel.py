#!/usr/bin/env python3
"""Đọc file Excel thiết kế (ユースケース記述 / 画面設計書 / 業務フロー) thành Markdown hoặc JSON.

Lấy được những thứ mà openpyxl đọc thường bỏ sót:
  - ô gộp (merged cell): ghi kèm vùng gộp, ví dụ B3[B3:H3]
  - chữ trong shape / text box / nhóm shape (流れ図, 画面遷移) — đọc thẳng từ drawing XML
  - ảnh chụp màn hình / mockup dán vào sheet (tuỳ chọn lưu ra thư mục để xem)
  - comment (メモ), hyperlink, sheet / hàng / cột bị ẩn
  - ô gạch ngang (取消線) và chữ đỏ (赤字) — thường là chỗ đã sửa / đã bỏ trong đặc tả

Cách dùng:
  python read_usecase_excel.py design.xlsx --list                 # chỉ liệt kê sheet + từ khoá
  python read_usecase_excel.py design.xlsx                         # toàn bộ, Markdown
  python read_usecase_excel.py design.xlsx --sheets 報告 ログイン   # lọc sheet theo chuỗi con
  python read_usecase_excel.py design.xlsx --images-dir out/img    # lưu ảnh nhúng để mở xem
  python read_usecase_excel.py design.xlsx --format json -o uc.json

Chỉ hỗ trợ .xlsx / .xlsm. File .xls cũ: chuyển trước bằng
  soffice --headless --convert-to xlsx file.xls
"""
from __future__ import annotations

import argparse
import json
import posixpath
import re
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

try:
    import openpyxl
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover
    sys.exit("Cần openpyxl: pip install openpyxl")

NS = {
    "main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
    "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
}
R_EMBED = "{%s}embed" % NS["r"]
R_ID = "{%s}id" % NS["r"]

# Từ khoá giúp định vị sheet / mục trong đặc tả use case tiếng Nhật
KEYWORDS = [
    "ユースケース", "アクター", "概要", "目的", "トリガー", "事前条件", "事後条件",
    "基本フロー", "代替フロー", "例外フロー", "業務ルール", "ビジネスルール",
    "画面", "入力", "項目", "必須", "桁", "メッセージ", "エラー", "権限", "ロール",
    "AI", "判定", "分類", "推定", "要約", "ステータス", "テーブル", "備考",
]


# ---------------------------------------------------------------- tiện ích

def clean(text) -> str:
    if text is None:
        return ""
    s = str(text).replace("\r\n", "\n").replace("\r", "\n").strip()
    return s.replace("\n", " ⏎ ")


def is_red(font) -> bool:
    try:
        rgb = font.color.rgb if font and font.color is not None else None
        return isinstance(rgb, str) and rgb.upper().endswith("FF0000")
    except Exception:
        return False


def keyword_hits(texts) -> list[str]:
    joined = "\n".join(texts)
    return [k for k in KEYWORDS if k in joined]


# ---------------------------------------------------------------- drawing XML

def _rels(zf: zipfile.ZipFile, part: str) -> dict[str, tuple[str, str]]:
    """part = 'xl/worksheets/sheet1.xml' → {rId: (type, target_abs)}"""
    d, f = posixpath.split(part)
    rel_path = posixpath.join(d, "_rels", f + ".rels")
    if rel_path not in zf.namelist():
        return {}
    root = ET.fromstring(zf.read(rel_path))
    out = {}
    for r in root.findall("rel:Relationship", NS):
        target = r.get("Target", "")
        if r.get("TargetMode") == "External":
            abs_t = target
        elif target.startswith("/"):
            abs_t = target.lstrip("/")
        else:
            abs_t = posixpath.normpath(posixpath.join(d, target))
        out[r.get("Id")] = (r.get("Type", ""), abs_t)
    return out


def sheet_parts(zf: zipfile.ZipFile) -> dict[str, str]:
    """{sheet name: 'xl/worksheets/sheetN.xml'}"""
    wb = ET.fromstring(zf.read("xl/workbook.xml"))
    rels = _rels(zf, "xl/workbook.xml")
    out = {}
    for s in wb.findall("main:sheets/main:sheet", NS):
        rid = s.get(R_ID)
        if rid in rels:
            out[s.get("name")] = rels[rid][1]
    return out


def _anchor_cell(anchor) -> str:
    frm = anchor.find("xdr:from", NS)
    if frm is None:
        return "?"
    try:
        col = int(frm.findtext("xdr:col", default="0", namespaces=NS))
        row = int(frm.findtext("xdr:row", default="0", namespaces=NS))
        return f"{get_column_letter(col + 1)}{row + 1}"
    except ValueError:
        return "?"


def _anchor_sort_key(cell: str):
    m = re.match(r"([A-Z]+)(\d+)", cell)
    if not m:
        return (10**9, 10**9)
    col = 0
    for ch in m.group(1):
        col = col * 26 + ord(ch) - 64
    return (int(m.group(2)), col)


def _shape_text(el) -> str:
    paras = []
    for p in el.iter("{%s}p" % NS["a"]):
        t = "".join(x.text or "" for x in p.iter("{%s}t" % NS["a"]))
        if t.strip():
            paras.append(t.strip())
    return " / ".join(paras)


def read_drawings(zf: zipfile.ZipFile, sheet_part: str, images_dir: Path | None, sheet_idx: int):
    shapes, images = [], []
    for rtype, target in _rels(zf, sheet_part).values():
        if not rtype.endswith("/drawing") or target not in zf.namelist():
            continue
        droot = ET.fromstring(zf.read(target))
        drels = _rels(zf, target)
        anchors = (droot.findall("xdr:twoCellAnchor", NS) + droot.findall("xdr:oneCellAnchor", NS)
                   + droot.findall("xdr:absoluteAnchor", NS))
        for anc in anchors:
            cell = _anchor_cell(anc)
            # shape có chữ (kể cả nằm trong nhóm)
            for sp in anc.iter("{%s}sp" % NS["xdr"]):
                nv = sp.find("xdr:nvSpPr/xdr:cNvPr", NS)
                name = nv.get("name", "") if nv is not None else ""
                txt = _shape_text(sp)
                if txt:
                    shapes.append({"anchor": cell, "name": name, "text": txt})
            n_conn = sum(1 for _ in anc.iter("{%s}cxnSp" % NS["xdr"]))
            if n_conn:
                shapes.append({"anchor": cell, "name": f"(đường nối x{n_conn})", "text": ""})
            for pic in anc.iter("{%s}pic" % NS["xdr"]):
                nv = pic.find("xdr:nvPicPr/xdr:cNvPr", NS)
                name = nv.get("name", "") if nv is not None else ""
                descr = nv.get("descr", "") if nv is not None else ""
                blip = pic.find(".//a:blip", NS)
                media = None
                if blip is not None and blip.get(R_EMBED) in drels:
                    media = drels[blip.get(R_EMBED)][1]
                saved = None
                if images_dir and media and media in zf.namelist():
                    images_dir.mkdir(parents=True, exist_ok=True)
                    ext = posixpath.splitext(media)[1] or ".bin"
                    safe = re.sub(r"[^\w\-]+", "_", name)[:40] or "image"
                    out = images_dir / f"s{sheet_idx:02d}_{cell}_{safe}{ext}"
                    out.write_bytes(zf.read(media))
                    saved = str(out)
                images.append({"anchor": cell, "name": name, "descr": descr,
                               "media": media, "saved": saved})
    shapes.sort(key=lambda s: _anchor_sort_key(s["anchor"]))
    images.sort(key=lambda s: _anchor_sort_key(s["anchor"]))
    return shapes, images


# ---------------------------------------------------------------- đọc sheet

def read_sheet(ws, ws_formula, max_cells: int):
    merged_tl = {}
    for rng in ws.merged_cells.ranges:
        merged_tl[rng.coord.split(":")[0]] = rng.coord
    hidden_rows = {r for r, d in ws.row_dimensions.items() if d.hidden}
    hidden_cols = {c for c, d in ws.column_dimensions.items() if d.hidden}

    rows, comments, links, count = [], [], [], 0
    truncated = False
    for row in ws.iter_rows():
        cells = []
        for c in row:
            v = c.value
            if v is None and ws_formula is not None:
                fv = ws_formula[c.coordinate].value
                if isinstance(fv, str) and fv.startswith("="):
                    v = f"{{công thức {fv}}}"
            if c.comment is not None:
                comments.append({"cell": c.coordinate, "text": clean(c.comment.text)})
            if c.hyperlink is not None and c.hyperlink.target:
                links.append({"cell": c.coordinate, "target": c.hyperlink.target})
            if v is None or (isinstance(v, str) and not v.strip()):
                continue
            if hasattr(v, "isoformat"):
                v = v.isoformat()
            flags = []
            if c.font is not None and c.font.strike:
                flags.append("取消線")
            if is_red(c.font):
                flags.append("赤字")
            if c.row in hidden_rows or c.column_letter in hidden_cols:
                flags.append("ẩn")
            cells.append({
                "cell": c.coordinate,
                "merged": merged_tl.get(c.coordinate),
                "value": clean(v),
                "flags": flags,
            })
            count += 1
        if cells:
            rows.append({"row": row[0].row, "cells": cells})
        if count >= max_cells:
            truncated = True
            break
    return rows, comments, links, count, truncated


def read_workbook(path: Path, sheet_filters, images_dir: Path | None, max_cells: int, list_only: bool):
    wb = openpyxl.load_workbook(path, data_only=True)
    try:
        wb_f = openpyxl.load_workbook(path, data_only=False)
    except Exception:
        wb_f = None
    zf = zipfile.ZipFile(path)
    parts = sheet_parts(zf)

    result = {"file": str(path), "sheets": []}
    for idx, ws in enumerate(wb.worksheets, start=1):
        if sheet_filters and not any(f in ws.title for f in sheet_filters):
            continue
        ws_f = wb_f[ws.title] if wb_f is not None else None
        rows, comments, links, count, truncated = read_sheet(ws, ws_f, max_cells)
        shapes, images = ([], [])
        if ws.title in parts:
            shapes, images = read_drawings(zf, parts[ws.title], None if list_only else images_dir, idx)
        texts = [c["value"] for r in rows for c in r["cells"]] + [s["text"] for s in shapes]
        result["sheets"].append({
            "index": idx,
            "name": ws.title,
            "state": ws.sheet_state,
            "dimension": ws.dimensions,
            "non_empty_cells": count,
            "truncated": truncated,
            "keywords": keyword_hits(texts),
            "rows": [] if list_only else rows,
            "shapes": [] if list_only else shapes,
            "images": images if not list_only else [{"anchor": i["anchor"]} for i in images],
            "comments": [] if list_only else comments,
            "links": [] if list_only else links,
        })
    return result


# ---------------------------------------------------------------- xuất

def to_markdown(res, list_only: bool) -> str:
    out = [f"# {Path(res['file']).name}", "", "## Danh sách sheet", "",
           "| # | Sheet | Trạng thái | Vùng | Ô có dữ liệu | Shape chữ | Ảnh | Từ khoá |",
           "|---|---|---|---|---|---|---|---|"]
    for s in res["sheets"]:
        n_shape = len([x for x in s["shapes"] if x["text"]]) if not list_only else "-"
        out.append(f"| {s['index']} | {s['name']} | {s['state']} | {s['dimension']} | "
                   f"{s['non_empty_cells']}{'+' if s['truncated'] else ''} | {n_shape} | "
                   f"{len(s['images'])} | {'、'.join(s['keywords'])} |")
    if list_only:
        return "\n".join(out) + "\n"

    for s in res["sheets"]:
        out += ["", f"## Sheet {s['index']}: {s['name']}" + ("  (ẨN)" if s["state"] != "visible" else ""), ""]
        if s["rows"]:
            out.append("### Ô (theo hàng; `ô[vùng gộp]` khi ô bị gộp)")
            out.append("")
            for r in s["rows"]:
                parts = []
                for c in r["cells"]:
                    ref = c["cell"] + (f"[{c['merged']}]" if c["merged"] else "")
                    flag = f" ‹{'・'.join(c['flags'])}›" if c["flags"] else ""
                    parts.append(f"{ref} {c['value']}{flag}")
                out.append(f"- R{r['row']}: " + " | ".join(parts))
            if s["truncated"]:
                out.append(f"- … (cắt bớt; dùng --max-cells lớn hơn để xem hết)")
        if s["shapes"]:
            out += ["", "### Shape / text box (theo vị trí neo)", ""]
            for sh in s["shapes"]:
                out.append(f"- {sh['anchor']} {sh['name']}: {sh['text']}" if sh["text"]
                           else f"- {sh['anchor']} {sh['name']}")
        if s["images"]:
            out += ["", "### Ảnh nhúng", ""]
            for im in s["images"]:
                where = f" → {im['saved']}" if im.get("saved") else ""
                desc = f" ({im['descr']})" if im.get("descr") else ""
                out.append(f"- {im['anchor']} {im['name']}{desc}{where}")
        if s["comments"]:
            out += ["", "### Comment (メモ)", ""]
            out += [f"- {c['cell']}: {c['text']}" for c in s["comments"]]
        if s["links"]:
            out += ["", "### Hyperlink", ""]
            out += [f"- {l['cell']}: {l['target']}" for l in s["links"]]
    return "\n".join(out) + "\n"


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("xlsx", type=Path)
    ap.add_argument("--sheets", nargs="*", default=[], help="chỉ đọc sheet có tên chứa chuỗi này")
    ap.add_argument("--list", action="store_true", help="chỉ liệt kê sheet")
    ap.add_argument("--images-dir", type=Path, help="lưu ảnh nhúng vào thư mục này")
    ap.add_argument("--max-cells", type=int, default=5000, help="giới hạn ô mỗi sheet (mặc định 5000)")
    ap.add_argument("--format", choices=["md", "json"], default="md")
    ap.add_argument("-o", "--out", type=Path)
    a = ap.parse_args()

    if a.xlsx.suffix.lower() not in {".xlsx", ".xlsm"}:
        sys.exit("Chỉ hỗ trợ .xlsx/.xlsm. Chuyển .xls bằng: soffice --headless --convert-to xlsx <file>")
    res = read_workbook(a.xlsx, a.sheets, a.images_dir, a.max_cells, a.list)
    text = json.dumps(res, ensure_ascii=False, indent=2) if a.format == "json" else to_markdown(res, a.list)
    if a.out:
        a.out.write_text(text, encoding="utf-8")
        print(f"Đã ghi {a.out}", file=sys.stderr)
    else:
        sys.stdout.write(text)


if __name__ == "__main__":
    main()
